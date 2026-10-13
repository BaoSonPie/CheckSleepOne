from docx import Document
import os
import time

# ==========================================
# CẤU HÌNH
# ==========================================

PROJECT_FOLDER = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DOCX = os.path.join(PROJECT_FOLDER, "CheckSleepOne.docx")


# ==========================================
# TẠO FILE WORD
# ==========================================

doc = Document()

doc.add_heading("CheckSleepOne - Python Source Code", level=1)

doc.add_paragraph(f"Thời gian xuất: {time.strftime('%Y-%m-%d %H:%M:%S')}")


# ==========================================
# DUYỆT TOÀN BỘ FILE PYTHON
# ==========================================

for root, dirs, files in os.walk(PROJECT_FOLDER):

    # Bỏ qua các thư mục không cần đưa vào Word
    dirs[:] = [
        d
        for d in dirs
        if d not in [".venv", "venv", "__pycache__", ".git", "runs", "dataset"]
    ]

    for file in sorted(files):

        if not file.endswith(".py"):
            continue

        # Không đưa chính file này vào Word
        if file == "export_py_to_word.py":
            continue

        path = os.path.join(root, file)

        # Đường dẫn tương đối
        relative_path = os.path.relpath(path, PROJECT_FOLDER)

        print(f"Đang ghi: {relative_path}")

        # Tiêu đề file
        doc.add_heading(relative_path, level=2)

        try:

            with open(path, "r", encoding="utf-8") as f:

                code = f.read()

            paragraph = doc.add_paragraph()
            run = paragraph.add_run(code)

            run.font.name = "Consolas"

        except Exception as e:

            doc.add_paragraph(f"Không thể đọc file: {e}")


# ==========================================
# LƯU WORD
# ==========================================

doc.save(OUTPUT_DOCX)

print()
print("==========================================")
print(" CHECKSLEEPONE CODE EXPORT")
print("==========================================")
print()
print("Đã xuất toàn bộ file Python vào:")
print(OUTPUT_DOCX)
print()
